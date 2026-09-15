from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw

from app.services.pdf_service import BRAND_LOGO, build_location_croquis, build_parcel_croquis, image_bytes_to_png

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "formulario_datos_tecnicos.docx"
CYAN = RGBColor(0x00, 0xA7, 0xD6)
NAVY = RGBColor(0x0B, 0x3A, 0x4A)
MUTED = RGBColor(0x4A, 0x6B, 0x75)
WHITE = RGBColor(255, 255, 255)
LINE = "C5E4EE"


def resolve_template_path() -> Path:
    """Plantilla activa: custom subida por admin, o builtin."""
    try:
        from app.services.template_service import get_active_template_path

        return get_active_template_path()
    except Exception:
        return TEMPLATE_PATH


def _render_with_docxtpl(template_path: Path, ctx: dict, images: dict[str, bytes]) -> bytes | None:
    """Rellena la plantilla .docx con docxtpl si tiene marcadores Jinja. None = usar builder interno."""
    try:
        from docxtpl import DocxTemplate, InlineImage
        from docx.shared import Inches as DocxInches
    except ImportError:
        return None
    if not template_path.exists():
        return None
    try:
        raw_probe = template_path.read_bytes()
        # Buscar marcadores solo en XML de Word. Revisar el zip entero da falsos
        # positivos (p. ej. bytes "{%" dentro de un PNG) y activa docxtpl sobre la
        # plantilla builtin → PDF con valores 0 y tablas vacías.
        import zipfile

        try:
            with zipfile.ZipFile(io.BytesIO(raw_probe)) as zf:
                xml_blob = b"".join(
                    zf.read(name)
                    for name in zf.namelist()
                    if name.startswith("word/") and name.endswith(".xml")
                )
        except zipfile.BadZipFile:
            return None

        from app.services.template_service import _custom_path

        is_custom = template_path.resolve() == _custom_path().resolve()
        has_jinja_control = b"{%" in xml_blob
        if not is_custom and not has_jinja_control:
            return None
        if b"{{" not in xml_blob and b"{%" not in xml_blob:
            return None
        tpl = DocxTemplate(str(template_path))
        render_ctx = dict(ctx)

        def _img(key: str, width_in: float = 1.6):
            data = images.get(key) or _placeholder_png(key)
            return InlineImage(tpl, io.BytesIO(_as_png(data)), width=DocxInches(width_in))

        render_ctx["croquis_predio"] = _img("croquis_predio", 2.4)
        render_ctx["croquis_ubicacion"] = _img("croquis_ubicacion", 2.4)
        render_ctx["foto_fachada1"] = _img("foto_fachada1", 1.5)
        render_ctx["foto_fachada2"] = _img("foto_fachada2", 1.5)
        render_ctx["foto_interior"] = _img("foto_interior", 1.5)
        # Alias cortos usados en plantillas de ejemplo
        for src, dst in (
            ("unit_number", "nro"),
            ("construction_year", "año"),
            ("kind", "bloque"),
            ("area", "sup"),
            ("floors_count", "pisos"),
            ("typology", "tipologia"),
            ("total_score", "puntaje"),
            ("modification_year", "modif"),
        ):
            for block in render_ctx.get("blocks") or []:
                if isinstance(block, dict) and src in block and dst not in block:
                    block[dst] = block[src]
        for row in render_ctx.get("all_chars") or []:
            if not isinstance(row, dict):
                continue
            row.setdefault("caracteristica", row.get("group"))
            row.setdefault("subtipo", row.get("option"))
            row.setdefault("ponderado", row.get("score"))
            row.setdefault("bloque", row.get("block"))
            row.setdefault("pct", row.get("pct"))
            row.setdefault("base", row.get("base"))

        tpl.render(render_ctx)
        buf = io.BytesIO()
        tpl.save(buf)
        return buf.getvalue()
    except Exception:
        return None


def build_formulario_docx(
    ctx: dict,
    images: dict[str, bytes],
    *,
    prefer_template: bool = True,
) -> bytes:
    """Arma el Word: plantilla activa (docxtpl) si aplica; si no, layout interno."""
    if prefer_template:
        filled = _render_with_docxtpl(resolve_template_path(), ctx, images)
        if filled:
            return filled
    return _build_formulario_docx_programmatic(ctx, images)


def _qr_payload(cadastral: str, property_number: str, pmc: str) -> str:
    return (
        "GAMC · Avalúo catastral\n"
        f"Código catastral: {cadastral or '—'}\n"
        f"Inmueble: {property_number or '—'}\n"
        f"PMC: {pmc or '—'}"
    )


def build_header_qr(cadastral: str | None, property_number: str | None, pmc: str | None) -> bytes:
    """QR compacto (navy sobre blanco) con ficha catastral para el encabezado."""
    import qrcode
    from qrcode.constants import ERROR_CORRECT_M

    qr = qrcode.QRCode(
        version=None,
        error_correction=ERROR_CORRECT_M,
        box_size=4,
        border=1,
    )
    qr.add_data(
        _qr_payload(
            str(cadastral).strip() if cadastral else "—",
            str(property_number).strip() if property_number else "—",
            str(pmc).strip() if pmc else "—",
        )
    )
    qr.make(fit=True)
    img = qr.make_image(fill_color=(11, 58, 74), back_color=(255, 255, 255)).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _money(n: float | None) -> str:
    return f"{float(n or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _num(n: float | None, decimals: int = 2) -> str:
    return f"{float(n or 0):.{decimals}f}"


def _dash(v) -> str:
    if v is None:
        return "—"
    s = str(v).strip()
    return s if s else "—"


def _placeholder_png(label: str, w: int = 420, h: int = 280) -> bytes:
    img = Image.new("RGB", (w, h), (244, 248, 250))
    d = ImageDraw.Draw(img)
    d.rectangle((1, 1, w - 2, h - 2), outline=(0, 167, 214))
    d.text((max(12, w // 2 - min(len(label) * 3, 90)), h // 2 - 6), label, fill=(74, 107, 117))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _set_run(run, *, size=9, bold=False, color=None, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def _p(cell, text, *, size=8, bold=False, color=None, align="left"):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    _set_run(run, size=size, bold=bold, color=color)
    return p


def _shade(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _borders(cell, color=LINE, sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _no_border_table(table):
    tblPr = table._tbl.tblPr if table._tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _heading_bar(doc, text: str):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    _shade(cell, "00A7D6")
    _p(cell, text, size=9, bold=True, color=WHITE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def _as_png(raw: bytes) -> bytes:
    try:
        return image_bytes_to_png(raw)
    except Exception:
        return raw


def _add_picture(cell, png: bytes, width_in: float = 1.35):
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cell.paragraphs[0].add_run()
    run.add_picture(io.BytesIO(_as_png(png)), width=Inches(width_in))


def _pick_photos(photo_paths: list[tuple[str, Path]]) -> tuple[bytes, bytes, bytes]:
    by_key: dict[str, bytes] = {}
    ordered: list[bytes] = []
    for name, path in photo_paths:
        try:
            raw = path.read_bytes()
        except OSError:
            continue
        if not raw:
            continue
        png = _as_png(raw)
        ordered.append(png)
        key = (name or "").lower()
        if "frente" in key or "front" in key or "fachada 1" in key:
            by_key.setdefault("f1", png)
        elif "lateral" in key or "lado" in key or "back" in key or "fondo" in key or "fachada 2" in key:
            by_key.setdefault("f2", png)
        elif "interior" in key or "perfil" in key or "vía" in key or "via" in key:
            by_key.setdefault("int", png)
    f1 = by_key.get("f1") or (ordered[0] if ordered else _placeholder_png("Fachada 1"))
    rest = [p for p in ordered if p is not f1]
    f2 = by_key.get("f2") or (rest[0] if rest else _placeholder_png("Fachada 2"))
    rest2 = [p for p in rest if p is not f2]
    interior = by_key.get("int") or (rest2[0] if rest2 else _placeholder_png("Interior"))
    return f1, f2, interior


def _new_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.2)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("GAMC · Formulario de actualización de datos técnicos")
    _set_run(r, size=7, color=MUTED)
    return doc


def _build_formulario_docx_programmatic(
    ctx: dict,
    images: dict[str, bytes],
) -> bytes:
    """Arma el Word con el layout del PDF legacy (invoicey)."""
    doc = _new_doc()

    head = doc.add_table(rows=1, cols=3)
    _no_border_table(head)
    head.autofit = False
    head.columns[0].width = Cm(2.3)
    head.columns[1].width = Cm(14.4)
    head.columns[2].width = Cm(2.7)
    logo_cell = head.cell(0, 0)
    logo_cell.width = Cm(2.3)
    if BRAND_LOGO.exists():
        logo_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_cell.paragraphs[0].add_run().add_picture(str(BRAND_LOGO), width=Inches(0.72))
    else:
        _p(logo_cell, "GAMC", size=10, bold=True, color=CYAN, align="center")

    title = head.cell(0, 1)
    title.width = Cm(14.4)
    title.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.paragraphs[0].add_run("GOBIERNO AUTÓNOMO MUNICIPAL DE COCHABAMBA")
    _set_run(r, size=11, bold=True, color=CYAN)
    p = title.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Cocha es progreso")
    _set_run(r, size=8, color=MUTED)
    p = title.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FORMULARIO PARA ACTUALIZACIÓN DE DATOS TÉCNICOS")
    _set_run(r, size=10, bold=True, color=NAVY)
    p = title.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DECLARACIÓN JURADA")
    _set_run(r, size=10, bold=True, color=NAVY)

    qr_cell = head.cell(0, 2)
    qr_cell.width = Cm(2.7)
    qr_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    qr_cell.paragraphs[0].clear()
    qr_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr_cell.paragraphs[0].paragraph_format.space_before = Pt(0)
    qr_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    qr_cell.paragraphs[0].add_run().add_picture(
        io.BytesIO(
            build_header_qr(
                ctx.get("cadastral_code"),
                ctx.get("property_number"),
                ctx.get("pmc"),
            )
        ),
        width=Cm(2.15),
    )
    cap = qr_cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(0)
    r = cap.add_run("Ficha catastral")
    _set_run(r, size=6, bold=True, color=CYAN)
    _borders(qr_cell, "C5E4EE", "4")

    _heading_bar(doc, "1.- Información del propietario")
    t = doc.add_table(rows=1, cols=1)
    _p(
        t.cell(0, 0),
        f"{ctx['owner_name']}    CI: {ctx['owner_document']}    {ctx['owner_city']}",
        size=9,
    )
    _borders(t.cell(0, 0))

    _heading_bar(doc, "2.- Información legal")
    legal = doc.add_table(rows=2, cols=3)
    pairs = [
        ("Matrícula", ctx["registry_matricula"]),
        ("Asiento", ctx["registry_asiento"]),
        ("Fecha DDRR", ctx["registry_ddr_date"]),
        ("Formulario N°", ctx["form_number"]),
        ("Fecha", ctx["form_date"]),
        ("Código", ctx["form_code"]),
    ]
    for i, (lab, val) in enumerate(pairs):
        cell = legal.cell(i // 3, i % 3)
        cell.paragraphs[0].clear()
        r = cell.paragraphs[0].add_run(f"{lab}: ")
        _set_run(r, size=8, color=MUTED)
        r = cell.paragraphs[0].add_run(str(val))
        _set_run(r, size=8, bold=True)
        _borders(cell)

    caps = doc.add_table(rows=1, cols=5)
    labels = ("Croquis del predio", "Fachada 1", "Fachada 2", "Interior", "Croquis de ubicación")
    keys = ("croquis_predio", "foto_fachada1", "foto_fachada2", "foto_interior", "croquis_ubicacion")
    for i, lab in enumerate(labels):
        _p(caps.cell(0, i), lab, size=7, bold=True, color=CYAN, align="center")
    imgs = doc.add_table(rows=1, cols=5)
    for i, key in enumerate(keys):
        cell = imgs.cell(0, i)
        _borders(cell, "00A7D6", "8")
        _add_picture(cell, images[key], width_in=1.32)

    _heading_bar(doc, "3.- Descripción del predio")
    desc = doc.add_table(rows=6, cols=2)
    desc_rows = [
        ("Calle / N°", f"{ctx['address']}  N° {ctx['door_number']}"),
        ("Edificio / Piso / Dpto.", f"{ctx['building_name']}  Piso: {ctx['floor_label']}  Dpto: {ctx['apartment_label']}"),
        ("Coordenadas", f"Ltd.: {ctx['latitude']}    Lgt.: {ctx['longitude']}"),
        ("Frente / Fondo / Superficie", f"{ctx['front_length']} m    {ctx['depth_length']} m    {ctx['approved_area']} m²"),
        ("Ubicación / Zona homogénea", f"{ctx['location_label']}    {ctx['zone_label']} ({ctx['zone_m2']} Bs/m²)"),
        ("Vía / Relieve / Forma", f"{ctx['road_label']}    {ctx['topo_label']}    {ctx['shape_label']}"),
    ]
    for i, (lab, val) in enumerate(desc_rows):
        _p(desc.cell(i, 0), lab, size=8, bold=True, color=CYAN)
        _p(desc.cell(i, 1), val, size=8)
        _borders(desc.cell(i, 0))
        _borders(desc.cell(i, 1))
        _shade(desc.cell(i, 0), "E8F7FB")
    svc = doc.add_paragraph()
    r = svc.add_run("Servicios: ")
    _set_run(r, size=8, bold=True, color=CYAN)
    r = svc.add_run(ctx["services"])
    _set_run(r, size=8)

    _heading_bar(doc, "4.- Valores calculados")
    vals = doc.add_table(rows=1, cols=4)
    for i, (lab, key) in enumerate(
        (
            ("Terreno (Bs)", "land_value"),
            ("Bloques (Bs)", "blocks_value"),
            ("Mejoras (Bs)", "improvements_value"),
            ("TOTAL (Bs)", "total_value"),
        )
    ):
        cell = vals.cell(0, i)
        cell.paragraphs[0].clear()
        r = cell.paragraphs[0].add_run(f"{lab}\n")
        _set_run(r, size=7, color=MUTED)
        r = cell.paragraphs[0].add_run(str(ctx.get(key) or "0,00"))
        _set_run(r, size=10, bold=True, color=NAVY)
        _borders(cell)
        if key == "total_value":
            _shade(cell, "E8F7FB")

    _heading_bar(doc, "5.- Características de la(s) construcciones")
    blocks = ctx["blocks"] or [
        {
            "unit_number": "—",
            "construction_year": "—",
            "kind": "—",
            "area": "—",
            "floors_count": "—",
            "typology": "—",
            "total_score": "—",
            "modification_year": "—",
        }
    ]
    bt = doc.add_table(rows=1 + len(blocks), cols=8)
    headers = ["Nro.", "Año", "Bloque", "Sup. cons.", "Pisos", "Tipología", "Puntaje", "Modif."]
    for i, h in enumerate(headers):
        _p(bt.cell(0, i), h, size=7, bold=True, color=WHITE, align="center")
        _shade(bt.cell(0, i), "00A7D6")
    keys_b = [
        "unit_number",
        "construction_year",
        "kind",
        "area",
        "floors_count",
        "typology",
        "total_score",
        "modification_year",
    ]
    for ri, b in enumerate(blocks, start=1):
        for ci, k in enumerate(keys_b):
            val = b.get(k, "—")
            if k == "area" and val not in ("—",):
                val = f"{val} m²"
            _p(bt.cell(ri, ci), str(val), size=7, align="center")
            _borders(bt.cell(ri, ci))

    decl = doc.add_paragraph()
    decl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = decl.add_run(
        "En mi calidad de sujeto pasivo y/o tercero responsable, declaro que la información "
        "proporcionada en la determinación del IPBI, fiel y exactamente refleja la verdad, por lo que "
        "juro a la exactitud de la presente declaración (Art.78,I, Ley 2492)."
    )
    _set_run(r, size=7, color=MUTED)
    decl.paragraph_format.space_before = Pt(8)

    signs = doc.add_table(rows=1, cols=2)
    _p(
        signs.cell(0, 0),
        f"_______________________\nFirma del propietario\n{ctx['owner_name']}\nCI: {ctx['owner_document']}",
        size=8,
        align="center",
    )
    _p(
        signs.cell(0, 1),
        f"_______________________\nFirma del profesional\nNro. registro: {ctx['professional_reg']}\n{ctx['professional_name']}",
        size=8,
        align="center",
    )

    _heading_bar(doc, "6.- Observaciones")
    obs = doc.add_table(rows=1, cols=1)
    _p(obs.cell(0, 0), ctx["observations"], size=8)
    _borders(obs.cell(0, 0))

    doc.add_page_break()
    head2 = doc.add_table(rows=1, cols=2)
    _no_border_table(head2)
    c0 = head2.cell(0, 0)
    r = c0.paragraphs[0].add_run("GOBIERNO AUTÓNOMO MUNICIPAL DE COCHABAMBA")
    _set_run(r, size=10, bold=True, color=CYAN)
    p = c0.add_paragraph()
    r = p.add_run("FORMULARIO PARA ACTUALIZACIÓN DE DATOS TÉCNICOS — Detalle de características")
    _set_run(r, size=9, bold=True, color=NAVY)
    _p(
        head2.cell(0, 1),
        f"Código catastral: {ctx['cadastral_code']}\n# Inmueble: {ctx['property_number']}",
        size=8,
        align="right",
    )

    chars = ctx["all_chars"]
    cht = doc.add_table(rows=1 + len(chars), cols=7)
    for i, h in enumerate(["Bloque", "#", "Característica", "Subtipo", "Puntaje base", "%", "Ponderado"]):
        _p(cht.cell(0, i), h, size=7, bold=True, color=WHITE, align="center")
        _shade(cht.cell(0, i), "0B3A4A")
    for ri, row in enumerate(chars, start=1):
        for ci, k in enumerate(["block", "n", "group", "option", "base", "pct", "score"]):
            _p(cht.cell(ri, ci), str(row.get(k, "—")), size=7)
            _borders(cht.cell(ri, ci))

    note = doc.add_paragraph()
    r = note.add_run(f"Observaciones: {ctx['observations']}")
    _set_run(r, size=8, color=MUTED)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def write_blank_template(path: Path | None = None) -> Path:
    """Plantilla vacía (mismos recuadros) para editar diseño de referencia."""
    dest = path or TEMPLATE_PATH
    dest.parent.mkdir(parents=True, exist_ok=True)
    empty = {k: f"{{{{{k}}}}}" for k in (
        "cadastral_code", "property_number", "pmc", "owner_name", "owner_document",
        "owner_city", "registry_matricula", "registry_asiento", "registry_ddr_date",
        "form_number", "form_date", "form_code", "address", "door_number",
        "building_name", "floor_label", "apartment_label", "latitude", "longitude",
        "front_length", "depth_length", "approved_area", "location_label",
        "zone_label", "zone_m2", "road_label", "topo_label", "shape_label",
        "services", "observations", "professional_name", "professional_reg",
    )}
    empty["blocks"] = [
        {
            "unit_number": "{{nro}}",
            "construction_year": "{{año}}",
            "kind": "{{bloque}}",
            "area": "{{sup}}",
            "floors_count": "{{pisos}}",
            "typology": "{{tipologia}}",
            "total_score": "{{puntaje}}",
            "modification_year": "{{modif}}",
        }
    ]
    empty["all_chars"] = [
        {
            "block": "{{bloque}}",
            "n": "1",
            "group": "{{caracteristica}}",
            "option": "{{subtipo}}",
            "base": "{{base}}",
            "pct": "{{pct}}",
            "score": "{{ponderado}}",
        }
    ]
    raw = build_formulario_docx(
        empty,
        {
            "croquis_predio": _placeholder_png("Croquis del predio"),
            "croquis_ubicacion": _placeholder_png("Croquis de ubicación"),
            "foto_fachada1": _placeholder_png("Fachada 1"),
            "foto_fachada2": _placeholder_png("Fachada 2"),
            "foto_interior": _placeholder_png("Interior"),
        },
        prefer_template=False,
    )
    dest.write_bytes(raw)
    return dest


def generate_appraisal_docx(
    data: dict,
    photo_paths: list[tuple[str, Path]] | None = None,
    *,
    professional_name: str | None = None,
    professional_reg: str | None = None,
    property_number: str | None = None,
    pmc: str | None = None,
    building_name: str | None = None,
    floor_label: str | None = None,
    apartment_label: str | None = None,
) -> bytes:
    rings = (data.get("map_meta") or {}).get("rings") or []
    parcel_png = build_parcel_croquis(
        float(data.get("front_length") or 0),
        float(data.get("depth_length") or 0),
        float(data.get("approved_area") or 0),
        data.get("cadastral_code"),
        rings=rings,
    )
    loc_png = build_location_croquis(data.get("latitude"), data.get("longitude"), rings=rings)
    f1, f2, interior = _pick_photos(photo_paths or [])

    blocks_ctx = []
    all_chars = []
    for b in data.get("blocks") or []:
        kind = "Mejora" if b.get("unit_kind") == "improvement" else "Bloque"
        typ = (b.get("typology") or {}).get("label") or ("Mejora" if b.get("unit_kind") == "improvement" else "—")
        title = f"{kind} {b.get('unit_number') or ''}".strip()
        blocks_ctx.append(
            {
                "unit_number": _dash(b.get("unit_number")),
                "construction_year": _dash(b.get("construction_year")),
                "kind": kind,
                "area": _num(b.get("area")),
                "floors_count": _dash(b.get("floors_count")),
                "typology": typ or "—",
                "total_score": _num(b.get("total_score")),
                "modification_year": _dash(b.get("modification_year")),
            }
        )
        chars = sorted(b.get("characteristics") or [], key=lambda x: x.get("group_sort", 0))
        if not chars:
            all_chars.append(
                {
                    "block": title or "—",
                    "n": "—",
                    "group": "Sin características",
                    "option": "—",
                    "base": "—",
                    "pct": "—",
                    "score": "—",
                }
            )
        for i, c in enumerate(chars, start=1):
            all_chars.append(
                {
                    "block": title or "—",
                    "n": str(i),
                    "group": c.get("group_name") or "—",
                    "option": c.get("option_label") or "—",
                    "base": _num(c.get("base_score")),
                    "pct": f"{float(c.get('percentage') or 0):.0f}",
                    "score": _num(c.get("score")),
                }
            )
    if not all_chars:
        all_chars.append(
            {
                "block": "—",
                "n": "—",
                "group": "—",
                "option": "—",
                "base": "—",
                "pct": "—",
                "score": "—",
            }
        )

    ctx = {
        "cadastral_code": _dash(data.get("cadastral_code")),
        "property_number": _dash(property_number),
        "pmc": _dash(pmc),
        "owner_name": _dash(data.get("owner_name")),
        "owner_document": _dash(data.get("owner_document")),
        "owner_city": "Cochabamba",
        "registry_matricula": _dash(data.get("registry_matricula")),
        "registry_asiento": _dash(data.get("registry_asiento")),
        "registry_ddr_date": _dash(data.get("registry_ddr_date")),
        "form_number": _dash(data.get("form_number")),
        "form_date": date.today().strftime("%d/%m/%Y"),
        "form_code": _dash(data.get("form_number")),
        "address": _dash(data.get("address")),
        "door_number": _dash(data.get("door_number")),
        "building_name": _dash(building_name),
        "floor_label": _dash(floor_label),
        "apartment_label": _dash(apartment_label),
        "latitude": _dash(data.get("latitude")),
        "longitude": _dash(data.get("longitude")),
        "front_length": _num(data.get("front_length")),
        "depth_length": _num(data.get("depth_length")),
        "approved_area": _num(data.get("approved_area")),
        "location_label": _dash(data.get("location_label")),
        "zone_label": _dash(data.get("zone_label")),
        "zone_m2": _num(data.get("zone_m2")),
        "road_label": _dash(data.get("road_label")),
        "topo_label": _dash(data.get("topo_label") or data.get("iprt_label")),
        "shape_label": _dash(data.get("shape_label")),
        "services": ", ".join(data.get("services") or []) or "—",
        "blocks": blocks_ctx,
        "all_chars": all_chars,
        "observations": _dash(data.get("observations")),
        "professional_name": _dash(professional_name),
        "professional_reg": _dash(professional_reg),
        "land_value": _money(data.get("land_value")),
        "blocks_value": _money(data.get("blocks_value")),
        "improvements_value": _money(data.get("improvements_value")),
        "total_value": _money(data.get("total_value")),
    }
    return build_formulario_docx(
        ctx,
        {
            "croquis_predio": parcel_png,
            "croquis_ubicacion": loc_png,
            "foto_fachada1": f1,
            "foto_fachada2": f2,
            "foto_interior": interior,
        },
    )


def convert_docx_to_pdf(docx_bytes: bytes) -> bytes:
    """Convierte el Word generado a PDF (no editable) con LibreOffice."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(
            "LibreOffice no está instalado en el servidor. Instale libreoffice-writer-nogui."
        )
    with tempfile.TemporaryDirectory(prefix="avaluo-docx-") as tmp:
        src = Path(tmp) / "formulario.docx"
        src.write_bytes(docx_bytes)
        env = os.environ.copy()
        env["HOME"] = tmp
        env["SAL_USE_VCLPLUGIN"] = "svp"
        proc = subprocess.run(
            [
                soffice,
                "--headless",
                "--norestore",
                "--invisible",
                "--nologo",
                "--nolockcheck",
                "--convert-to",
                "pdf:writer_pdf_Export",
                "--outdir",
                tmp,
                str(src),
            ],
            capture_output=True,
            timeout=90,
            env=env,
            cwd=tmp,
            check=False,
        )
        pdf = Path(tmp) / "formulario.pdf"
        if not pdf.exists() or pdf.stat().st_size < 100:
            err = (proc.stderr or proc.stdout or b"").decode("utf-8", "replace").strip()
            raise RuntimeError(err or "LibreOffice no produjo el PDF")
        return pdf.read_bytes()
